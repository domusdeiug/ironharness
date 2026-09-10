"""
`xlsx` profession — reads and edits spreadsheet (.xlsx) files.

Follows the same registration pattern used across app/professions/*/tools.py:
a `tools.py` that registers each tool on import via `register_tool`, with
narrow Pydantic args schemas and async handlers returning
`ToolExecutionResult`. Seed this profession's row via
`scripts/seed_professions.py`.

Tool names are namespaced (`xlsx.<name>`) per the registry's enforced-unique
global namespace — see `app/tools/registry.py`'s module docstring.

openpyxl's I/O is blocking; every handler runs the actual file work via
`asyncio.to_thread` so it doesn't stall the worker's event loop while
attempts on other requests are in flight.

error_type choices (read by ACTION's attempt-1 retry classification in
app/stages/action.py — _RETRYABLE_ERROR_TYPES vs _TERMINAL_ERROR_TYPES):
  - "not_found"      — file genuinely doesn't exist. No plausible arg
                        adjustment fixes this within the step; skip
                        straight to marking the step failed (ACTION will
                        escalate to replan if blocking).
  - unset (defaults to "transient", i.e. retryable) — malformed cell/range
    syntax, wrong sheet name, or a corrupt/unreadable workbook. These are
    exactly the "well-formed request, wrong details" cases attempts 2-3
    (cheap-model arg adjustment) exist for — marking them terminal would
    deny the retry ladder a chance to fix a typo'd sheet name or range.
"""

from __future__ import annotations

import asyncio
from enum import Enum
from typing import Any

import openpyxl
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from pydantic import BaseModel, Field

from app.infra.models import ToolExecutionResult
from app.tools.registry import register_tool


# ---------------------------------------------------------------------------
# read_range
# ---------------------------------------------------------------------------


class ReadRangeArgs(BaseModel):
    file_path: str
    cell_range: str = Field(
        ..., description="A1-style range, e.g. 'A1:D20'. A single cell like 'B2' is also valid."
    )
    sheet_name: str | None = Field(
        default=None, description="Defaults to the workbook's active sheet if omitted."
    )


def _read_range_sync(args: ReadRangeArgs) -> ToolExecutionResult:
    try:
        wb = openpyxl.load_workbook(args.file_path, data_only=True)
    except FileNotFoundError:
        return ToolExecutionResult(success=False, error=f"File not found: {args.file_path}", error_type="not_found")
    except Exception as exc:
        return ToolExecutionResult(success=False, error=f"Could not open workbook: {exc}")

    try:
        sheet = wb[args.sheet_name] if args.sheet_name else wb.active
    except KeyError:
        return ToolExecutionResult(success=False, error=f"Sheet not found: {args.sheet_name}")

    try:
        cells = sheet[args.cell_range]
    except Exception as exc:
        return ToolExecutionResult(success=False, error=f"Invalid cell range '{args.cell_range}': {exc}")

    if not isinstance(cells, tuple):
        cells = ((cells,),)
    elif cells and not isinstance(cells[0], tuple):
        cells = (cells,)

    values = [[c.value for c in row] for row in cells]
    return ToolExecutionResult(
        success=True,
        result={"sheet_name": sheet.title, "cell_range": args.cell_range, "values": values},
    )


async def _read_range(args: ReadRangeArgs) -> ToolExecutionResult:
    return await asyncio.to_thread(_read_range_sync, args)


# ---------------------------------------------------------------------------
# write_cells
# ---------------------------------------------------------------------------


class WriteCellsArgs(BaseModel):
    file_path: str
    updates: dict[str, str | float | int | bool | None] = Field(
        ..., description="Map of A1-style cell address -> value, e.g. {'B2': 42, 'C2': 'done'}."
    )
    sheet_name: str | None = None
    save_as: str | None = Field(
        default=None, description="If given, save to this path instead of overwriting file_path."
    )


def _write_cells_sync(args: WriteCellsArgs) -> ToolExecutionResult:
    try:
        wb = openpyxl.load_workbook(args.file_path)
    except FileNotFoundError:
        return ToolExecutionResult(success=False, error=f"File not found: {args.file_path}", error_type="not_found")
    except Exception as exc:
        return ToolExecutionResult(success=False, error=f"Could not open workbook: {exc}")

    try:
        sheet = wb[args.sheet_name] if args.sheet_name else wb.active
    except KeyError:
        return ToolExecutionResult(success=False, error=f"Sheet not found: {args.sheet_name}")

    try:
        for address, value in args.updates.items():
            sheet[address] = value
    except Exception as exc:
        return ToolExecutionResult(success=False, error=f"Invalid cell address in updates: {exc}")

    out_path = args.save_as or args.file_path
    try:
        wb.save(out_path)
    except Exception as exc:
        return ToolExecutionResult(success=False, error=f"Could not save workbook: {exc}")

    return ToolExecutionResult(
        success=True,
        result={"file_path": out_path, "sheet_name": sheet.title, "cells_written": len(args.updates)},
    )


async def _write_cells(args: WriteCellsArgs) -> ToolExecutionResult:
    return await asyncio.to_thread(_write_cells_sync, args)


# ---------------------------------------------------------------------------
# apply_formula
# ---------------------------------------------------------------------------


class ApplyFormulaArgs(BaseModel):
    file_path: str
    cell: str = Field(..., description="A1-style cell address the formula is written into, e.g. 'D10'.")
    formula: str = Field(..., description="Formula text including the leading '=', e.g. '=SUM(D2:D9)'.")
    sheet_name: str | None = None
    save_as: str | None = None


def _apply_formula_sync(args: ApplyFormulaArgs) -> ToolExecutionResult:
    if not args.formula.startswith("="):
        return ToolExecutionResult(success=False, error="formula must start with '='")

    try:
        wb = openpyxl.load_workbook(args.file_path)
    except FileNotFoundError:
        return ToolExecutionResult(success=False, error=f"File not found: {args.file_path}", error_type="not_found")
    except Exception as exc:
        return ToolExecutionResult(success=False, error=f"Could not open workbook: {exc}")

    try:
        sheet = wb[args.sheet_name] if args.sheet_name else wb.active
    except KeyError:
        return ToolExecutionResult(success=False, error=f"Sheet not found: {args.sheet_name}")

    try:
        sheet[args.cell] = args.formula
    except Exception as exc:
        return ToolExecutionResult(success=False, error=f"Invalid cell address '{args.cell}': {exc}")

    out_path = args.save_as or args.file_path
    try:
        wb.save(out_path)
    except Exception as exc:
        return ToolExecutionResult(success=False, error=f"Could not save workbook: {exc}")

    # openpyxl has no calc engine — the formula is written, not evaluated.
    # This is a code-checkable fact (not "success" in the intent-sense of
    # "and the computed value is correct"); Stage 3's evaluation-against-
    # intent step should treat this note as a signal, not just log it.
    return ToolExecutionResult(
        success=True,
        result={
            "file_path": out_path,
            "sheet_name": sheet.title,
            "cell": args.cell,
            "formula": args.formula,
            "note": "Formula written but not evaluated (openpyxl has no calc engine). "
            "Verify the computed value separately if a downstream step depends on it.",
        },
    )


async def _apply_formula(args: ApplyFormulaArgs) -> ToolExecutionResult:
    return await asyncio.to_thread(_apply_formula_sync, args)


# ---------------------------------------------------------------------------
# create_chart
# ---------------------------------------------------------------------------


class ChartType(str, Enum):
    bar = "bar"
    line = "line"
    pie = "pie"


class CreateChartArgs(BaseModel):
    file_path: str
    chart_type: ChartType
    data_range: str = Field(..., description="A1-style range containing the chart's data, e.g. 'B1:B10'.")
    categories_range: str | None = Field(
        default=None, description="A1-style range of category labels, e.g. 'A2:A10'."
    )
    title: str
    anchor_cell: str = Field(default="F2", description="Top-left cell the chart is anchored to.")
    sheet_name: str | None = None
    save_as: str | None = None


_CHART_CLASSES: dict[ChartType, Any] = {ChartType.bar: BarChart, ChartType.line: LineChart, ChartType.pie: PieChart}


def _create_chart_sync(args: CreateChartArgs) -> ToolExecutionResult:
    try:
        wb = openpyxl.load_workbook(args.file_path)
    except FileNotFoundError:
        return ToolExecutionResult(success=False, error=f"File not found: {args.file_path}", error_type="not_found")
    except Exception as exc:
        return ToolExecutionResult(success=False, error=f"Could not open workbook: {exc}")

    try:
        sheet = wb[args.sheet_name] if args.sheet_name else wb.active
    except KeyError:
        return ToolExecutionResult(success=False, error=f"Sheet not found: {args.sheet_name}")

    try:
        data_ref = Reference(sheet, range_string=f"{sheet.title}!{args.data_range}")
        chart = _CHART_CLASSES[args.chart_type]()
        chart.title = args.title
        chart.add_data(data_ref, titles_from_data=True)

        if args.categories_range:
            cats_ref = Reference(sheet, range_string=f"{sheet.title}!{args.categories_range}")
            chart.set_categories(cats_ref)

        sheet.add_chart(chart, args.anchor_cell)
    except Exception as exc:
        return ToolExecutionResult(success=False, error=f"Could not build chart: {exc}")

    out_path = args.save_as or args.file_path
    try:
        wb.save(out_path)
    except Exception as exc:
        return ToolExecutionResult(success=False, error=f"Could not save workbook: {exc}")

    return ToolExecutionResult(
        success=True,
        result={
            "file_path": out_path,
            "sheet_name": sheet.title,
            "chart_type": args.chart_type.value,
            "anchor_cell": args.anchor_cell,
        },
    )


async def _create_chart(args: CreateChartArgs) -> ToolExecutionResult:
    return await asyncio.to_thread(_create_chart_sync, args)


# ---------------------------------------------------------------------------
# Registration
# ---------------------------------------------------------------------------

register_tool(
    name="xlsx.read_range",
    args_schema=ReadRangeArgs,
    handler=_read_range,
    description="Read cell values from a range in an xlsx file.",
    skill_doc=(
        "xlsx.read_range(file_path, cell_range, sheet_name=None) -> "
        "{sheet_name, cell_range, values: 2D list}. sheet_name defaults to "
        "the active sheet. A missing sheet_name failure is retryable — "
        "inspect the workbook's actual sheet names and retry with the "
        "corrected one."
    ),
)

register_tool(
    name="xlsx.write_cells",
    args_schema=WriteCellsArgs,
    handler=_write_cells,
    description="Write literal values into one or more cells in an xlsx file.",
    skill_doc=(
        "xlsx.write_cells(file_path, updates: {A1_address: value}, "
        "sheet_name=None, save_as=None) -> {file_path, sheet_name, "
        "cells_written}. A malformed A1 address in `updates` (e.g. 'B' "
        "instead of 'B2') is the most common failure and is retryable "
        "with a corrected address."
    ),
)

register_tool(
    name="xlsx.apply_formula",
    args_schema=ApplyFormulaArgs,
    handler=_apply_formula,
    description="Write a formula into a single cell in an xlsx file.",
    skill_doc=(
        "xlsx.apply_formula(file_path, cell, formula, sheet_name=None, "
        "save_as=None) -> writes `formula` (must start with '=') into "
        "`cell`. Does NOT evaluate the formula — openpyxl has no calc "
        "engine, so 'success' here means 'written', not 'computed "
        "correctly'."
    ),
)

register_tool(
    name="xlsx.create_chart",
    args_schema=CreateChartArgs,
    handler=_create_chart,
    description="Add a bar, line, or pie chart to a sheet, built from a data range.",
    skill_doc=(
        "xlsx.create_chart(file_path, chart_type: bar|line|pie, "
        "data_range, categories_range=None, title, anchor_cell='F2', "
        "sheet_name=None, save_as=None). Pass bare A1 ranges (e.g. "
        "'B1:B10') — this tool builds the sheet-qualified reference "
        "internally."
    ),
)
